from __future__ import annotations
import csv
import io
import json
import shutil
import sqlite3
import socket
import hashlib
import re
import zipfile
from datetime import datetime, timedelta
from pathlib import Path

from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Request, Response, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .audit import append_audit
from .auth import create_session, current_user, get_db, require_role
from .config import APP_VERSION, BASE_DIR, MAX_URNS, DB_PATH, PUBLIC_KEY_PATH, PRIVATE_KEY_PATH, PRIVATE_KEY_PASSWORD_PATH, DATA_DIR
from .db import init_db
from .models import Ballot, DevicePrintCommand, Election, ElectionArchive, ElectionState, Occurrence, Setting, Slate, SlateMember, Urn, UrnStatus, User, UserRole, Voter, VoterExceptionPermit, VoterStatus, VotingAuthorization, AnonymousVotingAuthorization, AuditEvent, LoginGuard, Session as LoginSession, UrnDeviceSession
from .reports import audit_log_pdf, final_result_pdf, occurrences_pdf, voter_list_pdf, zero_pdf
from .security import canonical_json, hash_device_secret, hash_password, password_is_acceptable, random_token, sha256_text, verify_password
from .tally_service import tally_closed_election, verify_final_tally_integrity
from .device_auth import authenticate_device_secret, authenticate_urn_request, clear_urn_session_cookie, create_urn_session, set_urn_session_cookie
from .integrity import audit_head, security_status, verify_audit_chain
from .internal_key import ensure_internal_keypair, load_internal_private_key
from .services import (
    DomainError, authorize_voter, cancel_authorization, cast_vote, close_election,
    get_election, open_election, seal_election, verify_sealed_integrity, election_config_payload,
    append_voter_amendment,
)

app = FastAPI(title="Urna Escolar", version=APP_VERSION, docs_url=None, redoc_url=None, openapi_url=None)
USER_STATIC_DIR = DATA_DIR / "user_static"
USER_STATIC_DIR.mkdir(parents=True, exist_ok=True)
app.mount("/static", StaticFiles(directory=BASE_DIR / "app" / "static"), name="static")
app.mount("/user-static", StaticFiles(directory=USER_STATIC_DIR), name="user-static")
templates = Jinja2Templates(directory=BASE_DIR / "app" / "templates")
templates.env.globals["app_version"] = APP_VERSION


@app.exception_handler(HTTPException)
async def http_error(request: Request, exc: HTTPException):
    if exc.status_code == 401 and request.method == "GET" and request.url.path.startswith(("/admin", "/mesario")):
        from .db import SessionLocal
        with SessionLocal() as db:
            exists = db.scalar(select(User.id).limit(1)) is not None
        return RedirectResponse("/login" if exists else "/setup", status_code=303)
    return JSONResponse({"detail": exc.detail}, status_code=exc.status_code, headers=exc.headers)


@app.middleware("http")
async def security_headers(request: Request, call_next):
    # Bloqueia POST/PUT/PATCH/DELETE disparados por uma origem web diferente.
    # Cookies SameSite=Strict continuam sendo a primeira barreira; esta é uma
    # segunda defesa para operações administrativas e da urna.
    if request.method in {"POST", "PUT", "PATCH", "DELETE"}:
        origin = request.headers.get("origin")
        if origin:
            from urllib.parse import urlparse
            parsed = urlparse(origin)
            if parsed.scheme != request.url.scheme or parsed.netloc != request.url.netloc:
                return JSONResponse({"detail": "Origem da requisição não permitida."}, status_code=403)
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Permissions-Policy"] = "microphone=(), geolocation=()"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; "
        "base-uri 'self'; frame-ancestors 'none'; object-src 'none'; "
        "img-src 'self' data: blob:; media-src 'self'; "
        "style-src 'self' 'unsafe-inline'; script-src 'self' 'unsafe-inline'; "
        "connect-src 'self' http://127.0.0.1:8765 http://localhost:8765"
    )
    if request.url.scheme == "https":
        response.headers["Strict-Transport-Security"] = "max-age=31536000"
    if request.url.path.startswith(("/admin", "/mesario", "/urna", "/api", "/reports")):
        response.headers["Cache-Control"] = "no-store"
    return response


@app.on_event("startup")
def startup():
    init_db()


def _lan_ipv4() -> str:
    """Best-effort LAN IPv4 for the address other computers should open."""
    candidates: list[str] = []
    try:
        sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        try:
            sock.connect(("10.255.255.255", 1))
            candidates.append(sock.getsockname()[0])
        finally:
            sock.close()
    except Exception:
        pass
    try:
        for info in socket.getaddrinfo(socket.gethostname(), None, socket.AF_INET):
            candidates.append(info[4][0])
    except Exception:
        pass
    for ip in candidates:
        if ip and not ip.startswith("127.") and not ip.startswith("169.254."):
            return ip
    return "127.0.0.1"


def _network_base_url() -> str:
    return f"https://{_lan_ipv4()}:8443"


def _final_sound_url() -> str:
    override = USER_STATIC_DIR / "sounds" / "final.mp3"
    return "/user-static/sounds/final.mp3" if override.exists() else "/static/sounds/final.mp3"


def domain_call(db: Session, fn, *args, **kwargs):
    try:
        result = fn(db, *args, **kwargs)
        db.commit()
        return result
    except DomainError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Conflito de votação: o eleitor ou a urna já está em uso.") from exc
    except Exception:
        db.rollback()
        raise


def _login_guard_key(request: Request, username: str) -> str:
    ip = request.client.host if request.client else "unknown"
    return sha256_text(f"{ip}|{username.strip().lower()}")


def _assert_login_allowed(db: Session, key: str) -> None:
    guard = db.get(LoginGuard, key)
    now = datetime.utcnow()
    if guard and guard.locked_until and guard.locked_until > now:
        seconds = max(1, int((guard.locked_until - now).total_seconds()))
        raise HTTPException(429, f"Muitas tentativas. Aguarde cerca de {seconds // 60 + 1} minuto(s).")


def _record_login_failure(db: Session, key: str) -> None:
    now = datetime.utcnow()
    guard = db.get(LoginGuard, key)
    if not guard:
        guard = LoginGuard(key=key, failed_count=0, window_started_at=now)
        db.add(guard)
    if not guard.window_started_at or now - guard.window_started_at > timedelta(minutes=10):
        guard.failed_count = 0
        guard.window_started_at = now
        guard.locked_until = None
    guard.failed_count += 1
    if guard.failed_count >= 5:
        guard.locked_until = now + timedelta(minutes=5)


def _clear_login_guard(db: Session, key: str) -> None:
    guard = db.get(LoginGuard, key)
    if guard:
        db.delete(guard)


@app.get("/", response_class=HTMLResponse)
def root(request: Request):
    return templates.TemplateResponse(request=request, name="home.html", context={"version": APP_VERSION})


def _require_local_setup(request: Request) -> None:
    if not request.client or request.client.host not in {"127.0.0.1", "::1", "testclient"}:
        raise HTTPException(403, "Crie o primeiro administrador no computador da Central, usando o endereço local.")


@app.get("/setup", response_class=HTMLResponse)
def setup_page(request: Request, db: Session = Depends(get_db)):
    _require_local_setup(request)
    exists = db.execute(select(User).limit(1)).scalar_one_or_none() is not None
    return templates.TemplateResponse(request=request, name="setup.html", context={"exists": exists})


@app.post("/setup")
def setup_admin(request: Request, username: str = Form(...), display_name: str = Form(...), password: str = Form(...), db: Session = Depends(get_db)):
    _require_local_setup(request)
    if db.execute(select(User).limit(1)).scalar_one_or_none():
        raise HTTPException(409, "O sistema já possui usuário.")
    username, display_name = username.strip(), display_name.strip()
    if not username or not display_name:
        raise HTTPException(400, "Informe usuário e nome completo.")
    if len(username) > 80 or len(display_name) > 140:
        raise HTTPException(400, "Usuário ou nome ultrapassa o tamanho permitido.")
    if not password_is_acceptable(password):
        raise HTTPException(400, "A senha deve ter pelo menos 10 caracteres.")
    user = User(username=username, display_name=display_name, password_hash=hash_password(password), role=UserRole.ADMIN.value)
    db.add(user)
    get_election(db)
    db.commit()
    return RedirectResponse("/login", status_code=303)


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return templates.TemplateResponse(request=request, name="login.html", context={})


@app.post("/login")
def login(request: Request, username: str = Form(...), password: str = Form(...), db: Session = Depends(get_db)):
    key = _login_guard_key(request, username)
    _assert_login_allowed(db, key)
    user = db.execute(select(User).where(User.username == username.strip())).scalar_one_or_none()
    if not user or not user.active or not verify_password(user.password_hash, password):
        _record_login_failure(db, key)
        db.commit()
        raise HTTPException(401, "Usuário ou senha inválidos")
    _clear_login_guard(db, key)
    token = create_session(db, user)
    append_audit(db, "LOGIN", user.username, {"role": user.role})
    db.commit()
    target = "/admin" if user.role == UserRole.ADMIN.value else "/mesario"
    response = RedirectResponse(target, status_code=303)
    response.set_cookie(
        "session_id", token, max_age=4 * 3600, httponly=True,
        secure=(request.url.scheme == "https"), samesite="strict", path="/"
    )
    return response


@app.post("/logout")
def logout(request: Request, db: Session = Depends(get_db)):
    raw = request.cookies.get("session_id")
    if raw:
        session = db.get(LoginSession, sha256_text(raw))
        if session:
            db.delete(session)
            db.commit()
    response = RedirectResponse("/login", status_code=303)
    response.delete_cookie("session_id", path="/")
    return response


@app.get("/admin", response_class=HTMLResponse)
def admin_page(request: Request, section: str = "overview", user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    election = get_election(db)
    prep_sections = {"overview", "election", "voters", "slates", "urns", "users", "reports", "logs", "diagnostics", "security", "history"}
    readonly_sections = {"overview", "reports", "logs", "diagnostics", "security", "history"}
    if election.state == ElectionState.OPEN.value:
        # Durante a votação o Admin continua em modo de consulta, com uma única
        # exceção controlada: inclusão unitária de eleitor mediante nova senha
        # e ocorrência automática registrada para auditoria.
        readonly_sections = set(readonly_sections) | {"voters"}
    allowed_sections = prep_sections if election.state == ElectionState.CONFIG.value else readonly_sections
    if section not in allowed_sections:
        section = "overview"
    voters = db.execute(select(Voter).order_by(Voter.name)).scalars().all()
    slates = db.execute(select(Slate).order_by(Slate.number)).scalars().all()
    urns = db.execute(select(Urn).order_by(Urn.code)).scalars().all()
    users = db.execute(select(User).order_by(User.username)).scalars().all()
    audit_events = db.execute(select(AuditEvent).order_by(AuditEvent.id.desc()).limit(200)).scalars().all()
    voted = sum(1 for v in voters if v.status == VoterStatus.VOTED.value)
    in_progress = sum(1 for v in voters if v.status == VoterStatus.IN_PROGRESS.value)
    preflight = _preflight(db)
    archives = db.execute(select(ElectionArchive).order_by(ElectionArchive.archived_at.desc())).scalars().all()
    sec = security_status(db)
    live_config_hash = sha256_text(canonical_json(election_config_payload(db)))
    sec["config_current"] = live_config_hash
    sec["config_ok"] = (not election.config_hash or election.config_hash == live_config_hash)
    return templates.TemplateResponse(request=request, name="admin.html", context={
        "user": user, "election": election, "voters": voters, "voted": voted, "in_progress": in_progress,
        "slates": slates, "urns": urns, "users": users, "audit_events": audit_events, "archives": archives,
        "preflight": preflight, "security": sec, "max_urns": MAX_URNS, "section": section, "app_version": APP_VERSION,
        "network_base_url": _network_base_url(), "final_sound_url": _final_sound_url(),
    })


@app.post("/admin/election")
def admin_election(
    name: str = Form(...), institution_name: str = Form(""), allow_blank: str | None = Form(None),
    voter_identification_enabled: str | None = Form(None), voter_lookup_mode: str = Form("SEARCH_LIST"),
    show_members_on_ballot: str | None = Form(None), voting_mode: str = Form("SELECTION"),
    sound_enabled: str | None = Form(None), paper_enabled: str | None = Form(None),
    paper_change_warning_enabled: str | None = Form(None), paper_change_warning_votes: int = Form(200),
    paper_show_institution: str | None = Form(None), paper_show_election_name: str | None = Form(None),
    paper_show_slate_number: str | None = Form(None), paper_show_slate_name: str | None = Form(None),
    paper_show_members: str | None = Form(None), paper_show_instruction: str | None = Form(None),
    zero_copies: int = Form(3), result_copies: int = Form(3), zero_signature_fields: str | None = Form(None),
    keyboard_confirm_key: str = Form("Enter"), keyboard_correct_key: str = Form("Backspace"),
    keyboard_blank_key: str = Form("b"), paper_instruction_text: str = Form("Dobre e coloque na urna."),
    user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db),
):
    election = get_election(db)
    if election.state != ElectionState.CONFIG.value:
        raise HTTPException(409, "As configurações administrativas ficam bloqueadas depois da lacração. A abertura e o encerramento são feitos pela Mesa Eleitoral.")
    captured_keys = [(keyboard_confirm_key or "Enter").strip(), (keyboard_correct_key or "Backspace").strip(), (keyboard_blank_key or "b").strip()]
    if len({k.lower() for k in captured_keys}) != 3:
        raise HTTPException(400, "CONFIRMA, CORRIGE e BRANCO precisam usar teclas diferentes.")
    if not name.strip() or len(name.strip()) > 160 or len(institution_name.strip()) > 180:
        raise HTTPException(400, "Informe um nome de eleição e instituição dentro dos limites permitidos.")
    if voting_mode not in {"NUMERIC", "SELECTION"}:
        raise HTTPException(400, "Modo de votação inválido.")
    if any(not key or len(key) > 40 for key in captured_keys) or len(paper_instruction_text.strip()) > 180:
        raise HTTPException(400, "Tecla ou instrução de impressão inválida.")
    election.name = name.strip()
    election.institution_name = institution_name.strip()
    election.allow_blank = allow_blank is not None
    election.voter_identification_enabled = voter_identification_enabled is not None
    lookup = (voter_lookup_mode or "SEARCH_LIST").upper()
    if lookup not in {"SEARCH", "SEARCH_LIST", "LIST_ONLY"}:
        raise HTTPException(400, "Forma de navegação de eleitor inválida.")
    election.voter_lookup_mode = lookup
    election.show_members_on_ballot = show_members_on_ballot is not None
    election.voting_mode = voting_mode
    election.sound_enabled = sound_enabled is not None
    election.paper_enabled = paper_enabled is not None
    election.paper_change_warning_enabled = paper_change_warning_enabled is not None
    election.paper_change_warning_votes = max(20, min(int(paper_change_warning_votes or 200), 5000))
    election.paper_show_institution = paper_show_institution is not None
    election.paper_show_election_name = paper_show_election_name is not None
    election.paper_show_slate_number = paper_show_slate_number is not None
    election.paper_show_slate_name = paper_show_slate_name is not None
    election.paper_show_members = paper_show_members is not None
    election.paper_show_instruction = paper_show_instruction is not None
    election.zero_copies = max(3, min(zero_copies, 20))
    election.result_copies = max(1, min(result_copies, 20))
    election.zero_signature_fields = zero_signature_fields is not None
    election.keyboard_confirm_key = (keyboard_confirm_key or "Enter").strip()
    election.keyboard_correct_key = (keyboard_correct_key or "Backspace").strip()
    election.keyboard_blank_key = (keyboard_blank_key or "b").strip()
    election.paper_instruction_text = (paper_instruction_text or "Dobre e coloque na urna.").strip()
    append_audit(db, "ELECTION_SETTINGS_UPDATED", user.username, {"state": election.state})
    db.commit()
    return RedirectResponse("/admin?section=election", status_code=303)


@app.post("/admin/users")
def create_user(username: str = Form(...), display_name: str = Form(...), password: str = Form(...), role: str = Form(UserRole.MESARIO.value), user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    election = get_election(db)
    if election.state != ElectionState.CONFIG.value:
        raise HTTPException(409, "Usuários da eleição só podem ser alterados antes da lacração.")
    role = (role or UserRole.MESARIO.value).upper()
    if role not in {UserRole.ADMIN.value, UserRole.MESARIO.value}:
        raise HTTPException(400, "Perfil de usuário inválido.")
    username, display_name = username.strip(), display_name.strip()
    if not username or not display_name:
        raise HTTPException(400, "Informe usuário e nome completo.")
    if len(username) > 80 or len(display_name) > 140:
        raise HTTPException(400, "Usuário ou nome ultrapassa o tamanho permitido.")
    if not password_is_acceptable(password):
        raise HTTPException(400, "A senha deve ter pelo menos 10 caracteres.")
    if db.execute(select(User).where(User.username == username)).scalar_one_or_none():
        raise HTTPException(409, "Usuário já existe")
    db.add(User(username=username, display_name=display_name, password_hash=hash_password(password), role=role))
    append_audit(db, "USER_CREATED", user.username, {"new_username": username, "role": role})
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Usuário já existe")
    return RedirectResponse("/admin?section=users", status_code=303)


@app.post("/admin/voters/manual")
def add_voter(enrollment: str = Form(...), name: str = Form(...), class_code: str = Form(...), shift: str = Form(...), user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    if get_election(db).state != ElectionState.CONFIG.value:
        raise HTTPException(409, "Eleitores só podem ser alterados antes da lacração.")
    enrollment, name, class_code, shift = enrollment.strip(), name.strip(), class_code.strip(), shift.strip()
    if not enrollment or not name or not class_code or not shift:
        raise HTTPException(400, "Preencha matrícula, nome, turma e turno.")
    if len(enrollment) > 32 or len(name) > 180 or len(class_code) > 20 or len(shift) > 30:
        raise HTTPException(400, "Um dos campos ultrapassa o tamanho permitido.")
    if db.execute(select(Voter).where(Voter.enrollment == enrollment)).scalar_one_or_none():
        raise HTTPException(409, "Matrícula já cadastrada")
    db.add(Voter(enrollment=enrollment, name=name, class_code=class_code, shift=shift))
    append_audit(db, "VOTER_ADDED", user.username, {"enrollment": enrollment})
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Matrícula já cadastrada")
    return RedirectResponse("/admin?section=voters", status_code=303)


def _rows_from_upload(filename: str, content: bytes):
    """Return normalized source rows.

    XLSX: every worksheet is treated as one class/turma. The worksheet title is
    authoritative and becomes class_code for every student on that sheet.
    CSV: keeps the traditional Matricula/Nome/Turma/Turno columns.
    """
    lower = filename.lower()
    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        return [(None, row) for row in reader]
    if lower.endswith(".xlsx"):
        # XLSX is a ZIP: a small uploaded file can expand into many gigabytes.
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                raise HTTPException(400, "A planilha descompactada excede o limite de 50 MB.")
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        output = []
        try:
            for ws in wb.worksheets:
                if ws.max_row and ws.max_row > 50001 or ws.max_column and ws.max_column > 100:
                    raise HTTPException(400, "Use até 50.000 linhas e 100 colunas por planilha.")
                rows = ws.iter_rows(values_only=True)
                header = next(rows, None)
                if header is None:
                    continue
                headers = [str(x).strip() if x is not None else "" for x in header]
                class_name = str(ws.title).strip()
                for values in rows:
                    if not any(v is not None and str(v).strip() for v in values):
                        continue
                    output.append((class_name, dict(zip(headers, values))))
                    if len(output) > 50000:
                        raise HTTPException(400, "A importação permite até 50.000 eleitores por arquivo.")
        finally:
            wb.close()
        return output
    raise HTTPException(400, "Envie CSV ou XLSX")


@app.post("/admin/voters/import")
async def import_voters(file: UploadFile = File(...), user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    if get_election(db).state != ElectionState.CONFIG.value:
        raise HTTPException(409, "Eleitores só podem ser importados antes da lacração.")
    content = await file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "Arquivo de eleitores excede o limite de 10 MB.")
    try:
        rows = _rows_from_upload(file.filename or "", content)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, "Arquivo de eleitores inválido. Confira o formato CSV UTF-8 ou XLSX.") from exc
    if len(rows) > 50000:
        raise HTTPException(400, "A importação permite até 50.000 eleitores por arquivo.")
    aliases = {
        "matricula": "enrollment", "matrícula": "enrollment", "enrollment": "enrollment",
        "nome": "name", "name": "name", "turma": "class_code", "class_code": "class_code",
        "turno": "shift", "shift": "shift",
    }
    count = 0
    known_enrollments = set(db.scalars(select(Voter.enrollment)).all())
    for sheet_class, row in rows:
        normalized = {}
        for k, v in row.items():
            key = aliases.get(str(k).strip().lower())
            if key:
                normalized[key] = "" if v is None else str(v).strip()
        if sheet_class:
            normalized["class_code"] = sheet_class
        if not normalized.get("shift"):
            normalized["shift"] = "Não informado"
        if not all(normalized.get(k) for k in ("enrollment", "name", "class_code")):
            continue
        limits = {"enrollment": 32, "name": 180, "class_code": 20, "shift": 30}
        if any(len(normalized[field]) > limit for field, limit in limits.items()):
            raise HTTPException(400, "Um campo de eleitor ultrapassa o tamanho permitido. Nenhum eleitor foi importado.")
        if normalized["enrollment"] not in known_enrollments:
            db.add(Voter(**normalized))
            known_enrollments.add(normalized["enrollment"])
            count += 1
    append_audit(db, "VOTERS_IMPORTED", user.username, {"count": count})
    db.commit()
    return RedirectResponse("/admin?section=voters", status_code=303)


def _voter_exception_permit(db: Session, ticket: str, user: User) -> VoterExceptionPermit:
    token_hash = sha256_text((ticket or "").strip())
    permit = db.execute(
        select(VoterExceptionPermit).where(
            VoterExceptionPermit.token_hash == token_hash,
            VoterExceptionPermit.admin_user_id == user.id,
        )
    ).scalar_one_or_none()
    if not permit:
        raise HTTPException(401, "Autorização excepcional inválida. Inicie novamente.")
    if permit.consumed_at is not None:
        raise HTTPException(409, "Esta autorização já foi utilizada para cadastrar um eleitor.")
    if permit.expires_at < datetime.utcnow():
        raise HTTPException(401, "Esta autorização expirou. Inicie novamente e registre uma nova ocorrência de autorização.")
    return permit


@app.post("/api/admin/voters/exception/prepare")
def admin_voter_exception_prepare(
    password: str = Form(...),
    user: User = Depends(require_role(UserRole.ADMIN.value)),
    db: Session = Depends(get_db),
):
    election = get_election(db)
    if election.state != ElectionState.OPEN.value:
        raise HTTPException(409, "A inclusão excepcional só fica disponível enquanto a votação está aberta.")
    if not verify_password(user.password_hash, password or ""):
        raise HTTPException(401, "Senha do Administrador inválida.")
    try:
        verify_sealed_integrity(db)
    except DomainError as exc:
        raise HTTPException(409, str(exc)) from exc

    # Uma nova solicitação invalida qualquer autorização anterior ainda não usada
    # do mesmo Administrador, impedindo dois formulários simultâneos.
    previous = db.execute(
        select(VoterExceptionPermit).where(
            VoterExceptionPermit.admin_user_id == user.id,
            VoterExceptionPermit.consumed_at.is_(None),
        )
    ).scalars().all()
    now = datetime.utcnow()
    for old in previous:
        old.consumed_at = now

    registered = len(db.execute(select(Voter)).scalars().all())
    occurrence = Occurrence(
        author=f"Administrador: {user.display_name}",
        description=(
            "Solicitação de inclusão excepcional de eleitor durante a votação. "
            "O Administrador reautenticou a operação com sua senha. "
            "Esta ocorrência autoriza o cadastro de exatamente 1 (um) eleitor "
            "ausente da relação original. O cadastro geral permanece lacrado."
        ),
    )
    db.add(occurrence)
    db.flush()

    ticket = random_token(32)
    reference = f"OC-{occurrence.id:06d}"
    authorization_ref = reference
    permit = VoterExceptionPermit(
        token_hash=sha256_text(ticket), admin_user_id=user.id, print_command_id=authorization_ref,
        expires_at=now + timedelta(minutes=10), printed_at=now,
    )
    db.add(permit)
    append_audit(db, "VOTER_EXCEPTION_OCCURRENCE_RECORDED", user.username, {
        "reference": reference,
        "occurrence_id": occurrence.id,
        "authorization_ref": authorization_ref,
        "registered_voters_before": registered,
        "expires_minutes": 10,
    })
    db.commit()
    return {
        "ok": True, "ticket": ticket, "reference": reference,
        "occurrence_id": occurrence.id, "status": "ready",
        "message": "Ocorrência registrada. Cadastro unitário liberado para 1 eleitor.",
    }


@app.get("/api/admin/voters/exception/status")
def admin_voter_exception_status(
    ticket: str, user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db),
):
    try:
        permit = _voter_exception_permit(db, ticket, user)
    except HTTPException as exc:
        if exc.status_code in (401, 409):
            return JSONResponse({"status": "invalid", "message": exc.detail}, status_code=exc.status_code)
        raise
    return {
        "status": "ready",
        "message": "Ocorrência registrada. Cadastro unitário liberado.",
        "expires_at": permit.expires_at.isoformat() + "Z",
    }


@app.post("/api/admin/voters/exception/add")
def admin_voter_exception_add(
    ticket: str = Form(...), enrollment: str = Form(...), name: str = Form(...),
    class_code: str = Form(...), shift: str = Form(...),
    user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db),
):
    election = get_election(db)
    if election.state != ElectionState.OPEN.value:
        raise HTTPException(409, "A votação não está aberta.")
    permit = _voter_exception_permit(db, ticket, user)
    if permit.printed_at is None:
        raise HTTPException(409, "A ocorrência de autorização ainda não foi confirmada pelo sistema.")
    fields = {
        "enrollment": enrollment.strip(), "name": name.strip(),
        "class_code": class_code.strip(), "shift": shift.strip(),
    }
    if not all(fields.values()):
        raise HTTPException(400, "Preencha matrícula, nome, turma e turno.")
    if len(fields["enrollment"]) > 32 or len(fields["name"]) > 180 or len(fields["class_code"]) > 20 or len(fields["shift"]) > 30:
        raise HTTPException(400, "Um dos campos ultrapassa o tamanho permitido.")
    if db.execute(select(Voter).where(Voter.enrollment == fields["enrollment"])).scalar_one_or_none():
        raise HTTPException(409, "Matrícula já cadastrada. Esta autorização continua disponível para correção.")

    try:
        verify_sealed_integrity(db)
        now = datetime.utcnow()
        voter = Voter(
            **fields, exceptional_addition=True, exceptional_added_at=now,
            exceptional_added_by=user.username,
        )
        db.add(voter)
        db.flush()
        amendment = append_voter_amendment(db, voter, user.username, permit.print_command_id)
        permit.consumed_at = now
        append_audit(db, "VOTER_EXCEPTION_ADDED", user.username, {
            "enrollment": voter.enrollment, "name": voter.name,
            "class_code": voter.class_code, "shift": voter.shift,
            "voter_id": voter.id, "amendment_id": amendment.id,
            "authorization_ref": permit.print_command_id,
        })
        # A configuração originalmente lacrada continua intacta; a emenda
        # excepcional possui cadeia própria e precisa conferir antes do commit.
        verify_sealed_integrity(db)
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Matrícula já cadastrada. Esta autorização continua disponível para correção.")
    except DomainError as exc:
        db.rollback()
        raise HTTPException(409, str(exc)) from exc
    return {
        "ok": True,
        "voter": {"enrollment": voter.enrollment, "name": voter.name, "class_code": voter.class_code, "shift": voter.shift},
        "message": "Eleitor cadastrado e sincronizado. A autorização foi encerrada.",
    }


@app.post("/admin/slates")
async def add_slate(number: int = Form(...), name: str = Form(...), members_json: str = Form("[]"), logo: UploadFile | None = File(None), user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    if get_election(db).state != ElectionState.CONFIG.value:
        raise HTTPException(409, "Chapas só podem ser alteradas antes da lacração.")
    if number < 1 or number > 99:
        raise HTTPException(400, "Número deve ser de 1 a 99")
    clean_name = name.strip()
    if not clean_name or len(clean_name) > 150:
        raise HTTPException(400, "Informe o nome da chapa.")
    if db.execute(select(Slate).where(Slate.number == number)).scalar_one_or_none():
        raise HTTPException(409, "Número de chapa já usado")
    try:
        members = json.loads(members_json or "[]")
    except json.JSONDecodeError:
        raise HTTPException(400, "Integrantes inválidos")
    if not isinstance(members, list) or len(members) > 100:
        raise HTTPException(400, "Informe uma lista de até 100 integrantes.")
    slate = Slate(number=number, name=clean_name)
    if logo and logo.filename:
        suffix = Path(logo.filename).suffix.lower()
        if suffix not in {".png", ".jpg", ".jpeg", ".webp"}:
            raise HTTPException(400, "Logo deve ser PNG, JPG ou WEBP")
        content = await logo.read(5 * 1024 * 1024 + 1)
        if not content or len(content) > 5 * 1024 * 1024:
            raise HTTPException(400, "A imagem da chapa deve ter no máximo 5 MB.")
        try:
            with Image.open(io.BytesIO(content)) as image:
                if image.format not in {"PNG", "JPEG", "WEBP"} or image.width * image.height > 16_000_000:
                    raise HTTPException(400, "Imagem inválida ou maior que 16 megapixels.")
                image.load()
                normalized_logo = io.BytesIO()
                image.convert("RGBA").save(normalized_logo, format="PNG")
                content = normalized_logo.getvalue()
        except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError) as exc:
            raise HTTPException(400, "O arquivo enviado não é uma imagem válida.") from exc
        upload_dir = USER_STATIC_DIR / "uploads" / "slates"
        upload_dir.mkdir(parents=True, exist_ok=True)
        file_name = f"chapa_{number}.png"
        (upload_dir / file_name).write_bytes(content)
        slate.logo_path = f"/user-static/uploads/slates/{file_name}"
    for i, item in enumerate(members):
        if isinstance(item, dict) and item.get("role") and item.get("name"):
            slate.members.append(SlateMember(role_name=str(item["role"]).strip()[:100], person_name=str(item["name"]).strip()[:160], display_order=i))
    db.add(slate)
    append_audit(db, "SLATE_ADDED", user.username, {"number": number, "name": clean_name})
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Número de chapa já usado")
    return RedirectResponse("/admin?section=slates", status_code=303)


@app.post("/admin/slates/{slate_id}/members")
def add_slate_member(slate_id: int, role_name: str = Form(...), person_name: str = Form(...), user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    if get_election(db).state != ElectionState.CONFIG.value:
        raise HTTPException(409, "Integrantes só podem ser alterados antes da lacração.")
    slate = db.get(Slate, slate_id)
    if not slate:
        raise HTTPException(404, "Chapa não encontrada")
    role_name, person_name = role_name.strip(), person_name.strip()
    if not role_name or not person_name:
        raise HTTPException(400, "Informe o cargo e o nome do integrante")
    order = max([m.display_order for m in slate.members], default=-1) + 1
    member = SlateMember(slate_id=slate.id, role_name=role_name[:100], person_name=person_name[:160], display_order=order)
    db.add(member)
    append_audit(db, "SLATE_MEMBER_ADDED", user.username, {"slate": slate.number, "role": role_name, "name": person_name})
    db.commit()
    return RedirectResponse("/admin?section=slates", status_code=303)


@app.post("/admin/slates/{slate_id}/members/{member_id}/delete")
def delete_slate_member(slate_id: int, member_id: int, user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    if get_election(db).state != ElectionState.CONFIG.value:
        raise HTTPException(409, "Integrantes só podem ser alterados antes da lacração.")
    slate = db.get(Slate, slate_id)
    member = db.get(SlateMember, member_id)
    if not slate or not member or member.slate_id != slate.id:
        raise HTTPException(404, "Integrante não encontrado")
    audit_data = {"slate": slate.number, "role": member.role_name, "name": member.person_name}
    db.delete(member)
    append_audit(db, "SLATE_MEMBER_REMOVED", user.username, audit_data)
    db.commit()
    return RedirectResponse("/admin?section=slates", status_code=303)


@app.post("/admin/urns")
def add_urn(code: str = Form(...), name: str = Form(...), device_secret: str = Form(...), user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    if get_election(db).state != ElectionState.CONFIG.value:
        raise HTTPException(409, "Urnas só podem ser configuradas antes da lacração.")
    count = len(db.execute(select(Urn)).scalars().all())
    if count >= MAX_URNS:
        raise HTTPException(409, "Limite de 3 urnas")
    code, name = code.strip().upper(), name.strip()
    if not re.fullmatch(r"[A-Z0-9_-]{1,20}", code) or not name or len(name) > 80:
        raise HTTPException(400, "Informe código e nome da urna.")
    if db.execute(select(Urn).where(Urn.code == code)).scalar_one_or_none():
        raise HTTPException(409, "Código de urna já cadastrado")
    if not password_is_acceptable(device_secret):
        raise HTTPException(400, "A senha da urna deve ter pelo menos 10 caracteres.")
    db.add(Urn(code=code, name=name, device_secret_hash=hash_device_secret(device_secret), status=UrnStatus.AVAILABLE.value))
    append_audit(db, "URN_ADDED", user.username, {"code": code})
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Código de urna já cadastrado")
    return RedirectResponse("/admin?section=urns", status_code=303)


@app.post("/admin/urns/{urn_id}/mode")
def set_urn_voting_mode(urn_id: int, voting_mode_override: str = Form("INHERIT"), user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    if get_election(db).state != ElectionState.CONFIG.value:
        raise HTTPException(409, "O modo individual da urna só pode ser alterado antes da lacração.")
    urn = db.get(Urn, urn_id)
    if not urn:
        raise HTTPException(404, "Urna não encontrada")
    mode = (voting_mode_override or "INHERIT").upper()
    if mode not in {"INHERIT", "NUMERIC", "SELECTION"}:
        raise HTTPException(400, "Modo de votação inválido")
    urn.voting_mode_override = mode
    append_audit(db, "URN_VOTING_MODE_UPDATED", user.username, {"urn": urn.code, "mode": mode})
    db.commit()
    return RedirectResponse("/admin?section=urns", status_code=303)


@app.post("/admin/seal")
def seal(user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    domain_call(db, seal_election, user.username)
    return RedirectResponse("/admin", status_code=303)



def _ensure_zero_snapshot(db: Session, actor: str, source: str = "pdf") -> dict:
    election = get_election(db)
    if election.state != ElectionState.SEALED.value:
        raise HTTPException(409, "A zerésima só pode ser emitida depois da lacração e antes da abertura da votação.")
    try:
        verify_sealed_integrity(db)
    except DomainError as exc:
        raise HTTPException(409, str(exc)) from exc
    ballot_count = len(db.execute(select(Ballot)).scalars().all())
    if ballot_count != 0:
        raise HTTPException(409, "A zerésima não pode ser emitida porque já existem votos registrados.")
    if election.zero_snapshot_json:
        return json.loads(election.zero_snapshot_json)
    slates = db.execute(select(Slate).where(Slate.active == True).order_by(Slate.number)).scalars().all()
    urns = db.execute(select(Urn).order_by(Urn.code)).scalars().all()
    voter_count = len(db.execute(select(Voter)).scalars().all())
    issued = datetime.utcnow()
    snapshot = {
        "issued_at": issued.isoformat(),
        "election": election.name,
        "institution": election.institution_name,
        "software_version": APP_VERSION,
        "election_uuid": election.election_uuid,
        "config_hash": election.config_hash,
        "software_hash": election.software_hash,
        "public_key_fingerprint": election.public_key_fingerprint,
        "audit_head_before_zero": audit_head(db),
        "registered_voters": voter_count if election.voter_identification_enabled else None,
        "voter_identification_enabled": bool(election.voter_identification_enabled),
        "ballot_count": 0,
        "allow_blank": bool(election.allow_blank),
        "slates": [{"number": s.number, "name": s.name, "votes": 0} for s in slates],
        "urns": [{"code": u.code, "name": u.name} for u in urns],
    }
    raw = json.dumps(snapshot, ensure_ascii=False, sort_keys=True)
    election.zero_snapshot_json = raw
    election.zero_snapshot_hash = sha256_text(raw)
    election.zero_issued_at = issued
    append_audit(db, "ZERO_ISSUED", actor, {"source": source, "snapshot_hash": election.zero_snapshot_hash})
    return snapshot


@app.get("/reports/zero.pdf")
def report_zero(user: User = Depends(current_user), db: Session = Depends(get_db)):
    election = get_election(db)
    if not election.zero_snapshot_json:
        raise HTTPException(409, "A zerésima ainda não foi emitida. Ela é gerada pela Mesa Eleitoral durante a preparação da abertura.")
    pdf = zero_pdf(db)
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": "inline; filename=zeresima.pdf"})


@app.get("/reports/voters.pdf")
def report_voters(status: str | None = None, signature: bool = True, sort: str = "class", shift: str | None = None, class_code: str | None = None, class_pages: bool = False, user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    pdf = voter_list_pdf(db, only_status=status, signature=signature, sort=sort, shift=shift, class_code=class_code, class_pages=class_pages)
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": "inline; filename=eleitores.pdf"})



@app.get("/reports/result.pdf")
def report_result(user: User = Depends(current_user), db: Session = Depends(get_db)):
    try:
        pdf = final_result_pdf(db)
    except ValueError as exc:
        raise HTTPException(409, str(exc))
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": "inline; filename=boletim_final.pdf"})


@app.get("/mesario", response_class=HTMLResponse)
def mesario_page(request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if user.role != UserRole.MESARIO.value:
        raise HTTPException(403, "A Mesa Eleitoral exige login de usuário MESÁRIO.")
    election = get_election(db)
    urns = db.execute(select(Urn).order_by(Urn.code)).scalars().all()
    total = len(db.execute(select(Voter)).scalars().all()) if election.voter_identification_enabled else 0
    voted = (len(db.execute(select(Voter).where(Voter.status == VoterStatus.VOTED.value)).scalars().all())
             if election.voter_identification_enabled else len(db.execute(select(Ballot)).scalars().all()))
    classes = [row[0] for row in db.execute(select(Voter.class_code).distinct().order_by(Voter.class_code)).all()] if election.voter_identification_enabled else []
    return templates.TemplateResponse(request=request, name="mesario.html", context={"user": user, "election": election, "urns": urns, "total": total, "voted": voted, "classes": classes, "app_version": APP_VERSION, "preflight": _preflight(db), "network_base_url": _network_base_url()})



def _require_station_user(user: User) -> None:
    if user.role != UserRole.MESARIO.value:
        raise HTTPException(403, "Somente um usuário MESÁRIO pode operar a Mesa Eleitoral.")


def _reauth_station_user(user: User, password: str) -> None:
    if not verify_password(user.password_hash, password or ""):
        raise HTTPException(401, "Senha do mesário inválida.")


def _ticket_setting(db: Session, key: str, value: str | None = None) -> str | None:
    row = db.get(Setting, key)
    if value is None:
        return row.value if row else None
    if row is None:
        db.add(Setting(key=key, value=value))
    else:
        row.value = value
    return value


def _clear_open_ticket(db: Session) -> None:
    for key in ("mesario_open_ticket_hash", "mesario_open_ticket_user", "mesario_open_ticket_expires", "mesario_open_desk_printed"):
        row = db.get(Setting, key)
        if row:
            db.delete(row)


def _validate_open_ticket(db: Session, ticket: str, user: User) -> None:
    expected = _ticket_setting(db, "mesario_open_ticket_hash")
    owner = _ticket_setting(db, "mesario_open_ticket_user")
    expires = _ticket_setting(db, "mesario_open_ticket_expires")
    if not expected or not ticket or sha256_text(ticket) != expected or owner != str(user.id):
        raise HTTPException(401, "Autorização de abertura inválida ou expirada.")
    try:
        expiry = datetime.fromisoformat(expires or "")
    except Exception:
        expiry = datetime.min
    if expiry < datetime.utcnow():
        _clear_open_ticket(db)
        db.commit()
        raise HTTPException(401, "A autorização de abertura expirou. Inicie novamente.")


def _device_document_statuses(db: Session, kind: str) -> list[dict]:
    urns = db.execute(select(Urn).order_by(Urn.code)).scalars().all()
    items = []
    for urn in urns:
        cmd = db.execute(
            select(DevicePrintCommand).where(DevicePrintCommand.urn_id == urn.id, DevicePrintCommand.kind == kind)
            .order_by(DevicePrintCommand.created_at.desc()).limit(1)
        ).scalar_one_or_none()
        state = "pending"
        if cmd and cmd.success is True:
            state = "ok"
        elif cmd and cmd.success is False:
            state = "failed"
        items.append({"code": urn.code, "name": urn.name, "state": state})
    return items


def _latest_device_document_status(db: Session, kind: str) -> tuple[int, int, int]:
    items = _device_document_statuses(db, kind)
    completed = sum(1 for item in items if item["state"] == "ok")
    failed = sum(1 for item in items if item["state"] == "failed")
    return len(items), completed, failed


@app.post("/api/mesario/open/prepare")
def mesario_open_prepare(
    password: str = Form(...), user: User = Depends(current_user), db: Session = Depends(get_db)
):
    _require_station_user(user)
    _reauth_station_user(user, password)
    election = get_election(db)
    if election.state != ElectionState.SEALED.value:
        raise HTTPException(409, "A eleição precisa estar lacrada para ser aberta pela mesa.")

    ignore = {"zero", "zero_printed"}
    blocking = [c for c in _preflight(db) if c["blocking"] and not c["ok"] and c["key"] not in ignore]
    if blocking:
        raise HTTPException(409, "A eleição ainda não está pronta: " + "; ".join(c["label"] for c in blocking))

    urns = db.execute(select(Urn).order_by(Urn.code)).scalars().all()
    missing = [u.name for u in urns if not u.printer_name]
    if missing:
        raise HTTPException(409, "Selecione a impressora destas urnas antes da abertura: " + ", ".join(missing))

    _ensure_zero_snapshot(db, user.username, "mesario_open")
    queued = _queue_device_document(db, "zero", lambda urn: _zero_thermal_text(db, urn), election.zero_copies)
    ticket = random_token(24)
    _ticket_setting(db, "mesario_open_ticket_hash", sha256_text(ticket))
    _ticket_setting(db, "mesario_open_ticket_user", str(user.id))
    _ticket_setting(db, "mesario_open_ticket_expires", (datetime.utcnow() + timedelta(minutes=5)).isoformat())
    append_audit(db, "OPENING_ZERO_QUEUED", user.username, {"copies": election.zero_copies, "urns": queued})
    db.commit()
    return {
        "ok": True, "ticket": ticket, "copies": election.zero_copies,
        "urns": [{"code": u.code, "name": u.name} for u in urns],
        "election_name": election.name, "integrity_ok": True,
    }


@app.get("/api/mesario/open/ready")
def mesario_open_ready(ticket: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    _require_station_user(user)
    _validate_open_ticket(db, ticket, user)
    items = _device_document_statuses(db, "zero")
    completed = sum(1 for item in items if item["state"] == "ok")
    failed = sum(1 for item in items if item["state"] == "failed")
    ready = bool(items) and completed == len(items) and failed == 0
    return {"ready": ready, "total": len(items), "completed": completed, "failed": failed, "urns": items}


@app.post("/api/mesario/open/commit")
def mesario_open_commit(ticket: str = Form(...), user: User = Depends(current_user), db: Session = Depends(get_db)):
    _require_station_user(user)
    _validate_open_ticket(db, ticket, user)
    total, completed, failed = _latest_device_document_status(db, "zero")
    if not total or completed != total or failed:
        raise HTTPException(409, "A zerésima ainda não foi confirmada em todas as urnas.")
    try:
        open_election(db, user.username)
        _clear_open_ticket(db)
        db.commit()
        return {"ok": True, "state": ElectionState.OPEN.value}
    except DomainError as exc:
        db.rollback()
        raise HTTPException(409, str(exc)) from exc


@app.post("/api/mesario/close")
def mesario_close(
    password: str = Form(...), confirm_close: str = Form(...),
    user: User = Depends(current_user), db: Session = Depends(get_db),
):
    _require_station_user(user)
    _reauth_station_user(user, password)
    if confirm_close != "yes":
        raise HTTPException(400, "Confirme o encerramento da votação.")
    election = get_election(db)
    if election.state != ElectionState.OPEN.value:
        raise HTTPException(409, "A votação não está aberta.")
    try:
        private_key_bytes, key_password = load_internal_private_key()
        close_election(db, user.username)
        result = tally_closed_election(db, private_key_bytes, key_password, user.username)
        queued = _queue_device_document(db, "result", lambda urn: _result_thermal_text(db, urn), election.result_copies)
        close_ticket = random_token(24)
        _ticket_setting(db, "mesario_close_ticket_hash", sha256_text(close_ticket))
        _ticket_setting(db, "mesario_close_ticket_user", str(user.id))
        _ticket_setting(db, "mesario_close_ticket_expires", (datetime.utcnow() + timedelta(minutes=15)).isoformat())
        append_audit(db, "FINAL_RESULT_AUTO_QUEUED", user.username, {"urns": queued, "copies": election.result_copies, "source": "mesario"})
        db.commit()
        return {
            "ok": True, "state": ElectionState.CLOSED.value, "result": result,
            "copies": election.result_copies, "print_queued": queued,
            "election_name": election.name, "ticket": close_ticket,
        }
    except FileNotFoundError as exc:
        db.rollback()
        raise HTTPException(409, str(exc)) from exc
    except DomainError as exc:
        db.rollback()
        raise HTTPException(409, str(exc)) from exc
    except Exception:
        db.rollback()
        raise


def _validate_close_ticket(db: Session, ticket: str, user: User) -> None:
    expected = _ticket_setting(db, "mesario_close_ticket_hash")
    owner = _ticket_setting(db, "mesario_close_ticket_user")
    expires = _ticket_setting(db, "mesario_close_ticket_expires")
    if not expected or not ticket or sha256_text(ticket) != expected or owner != str(user.id):
        raise HTTPException(401, "Autorização de encerramento inválida ou expirada.")
    try:
        expiry = datetime.fromisoformat(expires or "")
    except Exception:
        expiry = datetime.min
    if expiry < datetime.utcnow():
        raise HTTPException(401, "A autorização de encerramento expirou.")


@app.get("/api/mesario/close/ready")
def mesario_close_ready(ticket: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    _require_station_user(user)
    _validate_close_ticket(db, ticket, user)
    items = _device_document_statuses(db, "result")
    completed = sum(1 for item in items if item["state"] == "ok")
    failed = sum(1 for item in items if item["state"] == "failed")
    ready = bool(items) and completed == len(items) and failed == 0
    if ready:
        for key in ("mesario_close_ticket_hash", "mesario_close_ticket_user", "mesario_close_ticket_expires"):
            row = db.get(Setting, key)
            if row:
                db.delete(row)
        db.commit()
    return {"ready": ready, "total": len(items), "completed": completed, "failed": failed, "urns": items}


@app.get("/api/voters/search")
def voter_search(q: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    _require_station_user(user)
    election = get_election(db)
    if election.state != ElectionState.OPEN.value:
        raise HTTPException(409, "A votação não está aberta.")
    if not election.voter_identification_enabled:
        raise HTTPException(409, "A identificação de eleitores está desativada nesta eleição.")
    if election.voter_lookup_mode == "LIST_ONLY":
        raise HTTPException(409, "A busca por digitação está desativada. Use a lista por turma.")
    term = q.strip()
    voters = db.execute(
        select(Voter).where(or_(Voter.enrollment == term, Voter.name.ilike(f"%{term}%"))).order_by(Voter.name).limit(15)
    ).scalars().all()
    urn_map = {u.id: u for u in db.execute(select(Urn)).scalars().all()}
    return [{"enrollment": v.enrollment, "name": v.name, "class_code": v.class_code, "shift": v.shift, "status": v.status, "active_urn_id": v.active_urn_id, "active_urn_code": urn_map[v.active_urn_id].code if v.active_urn_id in urn_map else None, "active_urn_name": urn_map[v.active_urn_id].name if v.active_urn_id in urn_map else None} for v in voters]


@app.get("/api/voters/classes")
def voter_classes(user: User = Depends(current_user), db: Session = Depends(get_db)):
    _require_station_user(user)
    election = get_election(db)
    if not election.voter_identification_enabled:
        return {"enabled": False, "classes": []}
    rows = db.execute(select(Voter.class_code, func.count(Voter.id)).group_by(Voter.class_code).order_by(Voter.class_code)).all()
    return {"enabled": True, "mode": election.voter_lookup_mode, "classes": [{"name": c or "Sem turma", "count": int(n)} for c, n in rows]}


@app.get("/api/voters/list")
def voter_list(class_code: str = "ALL", user: User = Depends(current_user), db: Session = Depends(get_db)):
    _require_station_user(user)
    election = get_election(db)
    if election.state != ElectionState.OPEN.value:
        raise HTTPException(409, "A votação não está aberta.")
    if not election.voter_identification_enabled:
        return []
    if election.voter_lookup_mode not in {"SEARCH_LIST", "LIST_ONLY"}:
        raise HTTPException(409, "A lista por turma está desativada nesta eleição.")
    stmt = select(Voter)
    if class_code and class_code.upper() != "ALL":
        stmt = stmt.where(Voter.class_code == class_code)
    voters = db.execute(stmt.order_by(Voter.name)).scalars().all()
    urn_map = {u.id: u for u in db.execute(select(Urn)).scalars().all()}
    return [{"enrollment": v.enrollment, "name": v.name, "class_code": v.class_code, "shift": v.shift, "status": v.status,
             "active_urn_id": v.active_urn_id,
             "active_urn_code": urn_map[v.active_urn_id].code if v.active_urn_id in urn_map else None,
             "active_urn_name": urn_map[v.active_urn_id].name if v.active_urn_id in urn_map else None} for v in voters]


@app.post("/api/mesario/authorize")
def mesario_authorize(enrollment: str = Form(...), urn_code: str = Form(...), user: User = Depends(current_user), db: Session = Depends(get_db)):
    with deployment.write_lock:
        _require_station_user(user)
        domain_call(db, authorize_voter, enrollment, urn_code, user.username)
        return {"ok": True, "urn_code": urn_code}


@app.post("/api/mesario/cancel")
def mesario_cancel(enrollment: str = Form(...), user: User = Depends(current_user), db: Session = Depends(get_db)):
    with deployment.write_lock:
        _require_station_user(user)
        domain_call(db, cancel_authorization, enrollment, user.username)
        return {"ok": True}


@app.get("/urna/{urn_code}", response_class=HTMLResponse)
def urn_page(request: Request, urn_code: str, db: Session = Depends(get_db)):
    urn = db.execute(select(Urn).where(Urn.code == urn_code.upper())).scalar_one_or_none()
    if not urn:
        raise HTTPException(404, "Urna não configurada")
    election = get_election(db)
    slates = db.execute(select(Slate).where(Slate.active == True).order_by(Slate.number)).scalars().all()
    return templates.TemplateResponse(request=request, name="urna.html", context={"urn": urn, "election": election, "slates": slates, "app_version": APP_VERSION, "final_sound_url": _final_sound_url()})


def _authenticate_urn(db: Session, urn_code: str, request: Request, secret: str | None = None) -> Urn:
    return authenticate_urn_request(db, urn_code, request, secret)


@app.post("/api/urna/{urn_code}/login")
def urn_device_login(request: Request, urn_code: str, device_secret: str = Form(...), db: Session = Depends(get_db)):
    guard_key = _login_guard_key(request, f"urna:{urn_code.upper()}")
    _assert_login_allowed(db, guard_key)
    try:
        urn = authenticate_device_secret(db, urn_code, device_secret)
    except HTTPException:
        _record_login_failure(db, guard_key)
        db.commit()
        raise
    _clear_login_guard(db, guard_key)
    token = create_urn_session(db, urn)
    urn.last_seen_at = datetime.utcnow()
    append_audit(db, "URN_DEVICE_LOGIN", f"device:{urn.code}", {"urn": urn.code})
    db.commit()
    response = JSONResponse({"ok": True, "urn": urn.code})
    set_urn_session_cookie(response, request, urn.code, token)
    return response


@app.post("/api/urna/{urn_code}/logout")
def urn_device_logout(request: Request, urn_code: str, db: Session = Depends(get_db)):
    raw = request.cookies.get('urn_session_' + ''.join(ch.lower() for ch in urn_code if ch.isalnum() or ch in ('-', '_')))
    if raw:
        session = db.get(UrnDeviceSession, sha256_text(raw))
        if session:
            db.delete(session)
            db.commit()
    response = JSONResponse({"ok": True})
    clear_urn_session_cookie(response, urn_code)
    return response


@app.get("/api/urna/{urn_code}/integrity-check")
def urn_integrity_check(request: Request, urn_code: str, x_urn_secret: str | None = Header(default=None), db: Session = Depends(get_db)):
    urn = _authenticate_urn(db, urn_code, request, x_urn_secret)
    election = get_election(db)
    checks = []
    ok = True
    try:
        if election.state == ElectionState.CONFIG.value:
            checks.append({"label": "Configuração", "ok": True, "detail": "Eleição ainda em preparação"})
        else:
            verify_sealed_integrity(db)
            checks += [
                {"label": "Software", "ok": True, "detail": "Hash confere com a lacração"},
                {"label": "Configuração", "ok": True, "detail": "Configuração lacrada íntegra"},
                {"label": "Chave interna", "ok": True, "detail": "Componente público e fingerprint correspondentes"},
                {"label": "Log de auditoria", "ok": True, "detail": "Cadeia de hashes íntegra"},
            ]
    except DomainError as exc:
        ok = False
        checks.append({"label": "Integridade", "ok": False, "detail": str(exc)})
    https_ok = request.url.scheme == "https"
    checks.append({"label": "HTTPS", "ok": https_ok, "detail": "Conexão HTTPS" if https_ok else "Conexão sem HTTPS"})
    ok = bool(ok and https_ok)
    append_audit(db, "URN_INTEGRITY_OK" if ok else "URN_INTEGRITY_BLOCKED", f"device:{urn.code}", {"urn": urn.code, "state": election.state})
    db.commit()
    return {"ok": ok, "checks": checks, "state": election.state}


@app.get("/api/urna/{urn_code}/authorization")
def urn_authorization(request: Request, urn_code: str, x_urn_secret: str | None = Header(default=None), db: Session = Depends(get_db)):
    urn = _authenticate_urn(db, urn_code, request, x_urn_secret)
    from .models import VotingAuthorization
    election = get_election(db)
    effective_mode = urn.voting_mode_override if (urn.voting_mode_override and urn.voting_mode_override != "INHERIT") else election.voting_mode
    if not election.voter_identification_enabled:
        auth = db.execute(select(AnonymousVotingAuthorization).where(AnonymousVotingAuthorization.urn_id == urn.id)).scalar_one_or_none()
        urn.last_seen_at = datetime.utcnow()
        db.commit()
        operational = {
            "election_state": election.state,
            "voting_mode": effective_mode,
            "sound_enabled": bool(election.sound_enabled),
            "keyboard_confirm_key": election.keyboard_confirm_key,
            "keyboard_correct_key": election.keyboard_correct_key,
            "keyboard_blank_key": election.keyboard_blank_key,
            "anonymous_mode": True,
        }
        return {"authorized": bool(auth), "token": auth.urn_token if auth else None, **operational}
    auth = db.execute(select(VotingAuthorization).where(VotingAuthorization.urn_id == urn.id)).scalar_one_or_none()
    db.commit()
    operational = {
        "election_state": election.state,
        "voting_mode": effective_mode,
        "sound_enabled": bool(election.sound_enabled),
        "keyboard_confirm_key": election.keyboard_confirm_key,
        "keyboard_correct_key": election.keyboard_correct_key,
        "keyboard_blank_key": election.keyboard_blank_key,
        "anonymous_mode": False,
    }
    if not auth:
        return {"authorized": False, **operational}
    return {"authorized": True, "token": auth.urn_token, **operational}


@app.post("/api/urna/{urn_code}/anonymous/start")
def urn_anonymous_start(request: Request, urn_code: str, x_urn_secret: str | None = Header(default=None), db: Session = Depends(get_db)):
    with deployment.write_lock:
        urn = _authenticate_urn(db, urn_code, request, x_urn_secret)
        election = get_election(db)
        if election.state != ElectionState.OPEN.value:
            raise HTTPException(409, "A votação não está aberta.")
        if election.voter_identification_enabled:
            raise HTTPException(409, "Esta eleição exige identificação de eleitor pela Mesa.")
        existing = db.execute(select(AnonymousVotingAuthorization).where(AnonymousVotingAuthorization.urn_id == urn.id).with_for_update()).scalar_one_or_none()
        if existing:
            return {"ok": True, "token": existing.urn_token}
        try:
            deployment.require_ready(urn, election)
        except DomainError as exc:
            raise HTTPException(409, str(exc)) from exc
        deployment.clear_previous_vote(db, urn)
        token = random_token(24)
        db.add(AnonymousVotingAuthorization(urn_token=token, urn_id=urn.id))
        urn.status = UrnStatus.IN_USE.value
        db.commit()
        return {"ok": True, "token": token}


@app.post("/api/urna/{urn_code}/vote")
def urn_vote(request: Request, urn_code: str, token: str = Form(...), choice_type: str = Form(...), slate_number: int | None = Form(None), x_urn_secret: str | None = Header(default=None), db: Session = Depends(get_db)):
    # Serialize read/check/write on SQLite. A replay returns the stored receipt,
    # never invokes cast_vote or queues a second print.
    with deployment.write_lock:
        urn = _authenticate_urn(db, urn_code, request, x_urn_secret)
        if urn.managed_print:
            receipt = deployment.vote_receipt(db, urn, token)
            if receipt:
                db.commit()
                return receipt
        choice = {"type": choice_type}
        if choice_type == "slate":
            choice["number"] = slate_number
        try:
            ballot_id = cast_vote(db, urn.code, token, choice)
            election = get_election(db)
            slate = db.execute(select(Slate).where(Slate.number == int(slate_number))).scalar_one() if choice_type == "slate" else None
            paper = {
                "enabled": election.paper_enabled and urn.print_enabled,
                "institution": election.institution_name if election.paper_show_institution else None,
                "election": election.name if election.paper_show_election_name else None,
                "number": slate.number if slate and election.paper_show_slate_number else None,
                "slate": slate.name if slate and election.paper_show_slate_name else ("BRANCO" if choice_type == "blank" else None),
                "members": [],
                "instruction": election.paper_instruction_text if election.paper_show_instruction else None,
            }
            if urn.managed_print:
                deployment.record_vote(db, urn, token, paper)
            db.commit()
            if urn.managed_print:
                return deployment.vote_receipt(db, urn, token)
            return {"ok": True, "ballot_id": ballot_id, "paper": paper}
        except DomainError as exc:
            db.rollback()
            raise HTTPException(409, str(exc)) from exc
        except Exception:
            db.rollback()
            raise


@app.post("/api/urna/{urn_code}/vote-status")
def urn_vote_status(request: Request, urn_code: str, token: str = Form(...), db: Session = Depends(get_db)):
    with deployment.write_lock:
        urn = _authenticate_urn(db, urn_code, request)
        receipt = deployment.vote_receipt(db, urn, token)
        db.commit()
        return receipt or {"recorded": False}


@app.post("/api/mesario/print-problem")
def mesario_print_problem(urn_code: str = Form(...), user: User = Depends(current_user), db: Session = Depends(get_db)):
    with deployment.write_lock:
        _require_station_user(user)
        urn = db.execute(select(Urn).where(Urn.code == urn_code.upper())).scalar_one_or_none()
        if not urn or not urn.pending_reprint_ballot_id:
            raise HTTPException(409, "Não há voto recente disponível para reimpressão nesta urna.")
        urn.status = UrnStatus.PRINT_ERROR.value
        append_audit(db, "PRINT_PROBLEM_REPORTED", user.username, {"urn": urn.code})
        db.commit()
        return {"ok": True}


@app.post("/api/mesario/reprint")
def mesario_reprint(urn_code: str = Form(...), user: User = Depends(current_user), db: Session = Depends(get_db)):
    with deployment.write_lock:
        _require_station_user(user)
        urn = db.execute(select(Urn).where(Urn.code == urn_code.upper())).scalar_one_or_none()
        if not urn or urn.status != UrnStatus.PRINT_ERROR.value or not urn.pending_reprint_ballot_id:
            raise HTTPException(409, "A urna não possui uma reimpressão pendente.")
        if urn.managed_print:
            deployment.reprint_vote(db, urn)
        else:
            urn.reprint_authorized = True
        append_audit(db, "REPRINT_AUTHORIZED", user.username, {"urn": urn.code})
        db.commit()
        return {"ok": True}


@app.get("/api/mesario/status")
def mesario_status(user: User = Depends(current_user), db: Session = Depends(get_db)):
    _require_station_user(user)
    election = get_election(db)
    total = len(db.execute(select(Voter)).scalars().all()) if election.voter_identification_enabled else None
    voted = (len(db.execute(select(Voter).where(Voter.status == VoterStatus.VOTED.value)).scalars().all())
             if election.voter_identification_enabled else int(db.execute(select(func.count(Ballot.id))).scalar_one() or 0))
    urns = db.execute(select(Urn).order_by(Urn.code)).scalars().all()

    paper_alerts = []
    threshold = max(20, min(int(election.paper_change_warning_votes or 200), 5000))
    if election.state == ElectionState.OPEN.value and election.paper_enabled and election.paper_change_warning_enabled:
        for urn in urns:
            printed_votes = int(db.execute(select(func.count(Ballot.id)).where(Ballot.urn_code == urn.code)).scalar_one() or 0)
            setting_key = f"paper_change_ack:{election.election_uuid}:{urn.code}"
            setting = db.get(Setting, setting_key)
            try:
                acknowledged_at = int(setting.value) if setting else 0
            except (TypeError, ValueError):
                acknowledged_at = 0
            since_change = max(0, printed_votes - acknowledged_at)
            if since_change >= threshold:
                paper_alerts.append({
                    "urn_code": urn.code, "urn_name": urn.name, "threshold": threshold,
                    "votes_since_change": since_change, "printed_votes": printed_votes,
                })

    return {
        "total": total, "voted": voted, "remaining": (max(0, total-voted) if total is not None else None), "state": election.state,
        "voter_identification_enabled": bool(election.voter_identification_enabled),
        "voter_lookup_mode": election.voter_lookup_mode,
        "preflight": _preflight(db),
        "paper_change": {
            "enabled": bool(election.paper_enabled and election.paper_change_warning_enabled),
            "threshold": threshold, "alerts": paper_alerts,
        },
        "urns": [{"code": u.code, "name": u.name, "status": u.status, "last_vote_at": u.last_vote_at.isoformat() if u.last_vote_at else None,
                  "can_report_print": bool(u.pending_reprint_ballot_id), "reprint_requested": u.reprint_authorized} for u in urns]
    }


@app.post("/api/mesario/paper-change/confirm")
def mesario_confirm_paper_change(urn_code: str = Form(...), user: User = Depends(current_user), db: Session = Depends(get_db)):
    _require_station_user(user)
    election = get_election(db)
    if election.state != ElectionState.OPEN.value:
        raise HTTPException(409, "A troca de papel só pode ser confirmada durante a votação aberta.")
    if not election.paper_enabled or not election.paper_change_warning_enabled:
        raise HTTPException(409, "O aviso de troca de papel não está habilitado nesta eleição.")
    urn = db.execute(select(Urn).where(Urn.code == urn_code)).scalar_one_or_none()
    if urn is None:
        raise HTTPException(404, "Urna não encontrada.")
    printed_votes = int(db.execute(select(func.count(Ballot.id)).where(Ballot.urn_code == urn.code)).scalar_one() or 0)
    setting_key = f"paper_change_ack:{election.election_uuid}:{urn.code}"
    setting = db.get(Setting, setting_key)
    if setting is None:
        db.add(Setting(key=setting_key, value=str(printed_votes)))
    else:
        setting.value = str(printed_votes)
    append_audit(db, "PAPER_ROLL_CHANGED", user.username, {"urn": urn.code, "printed_votes": printed_votes})
    db.commit()
    return {"ok": True, "urn_code": urn.code, "acknowledged_at": printed_votes}


@app.get("/api/urna/{urn_code}/reprint")
def urn_reprint_status(request: Request, urn_code: str, x_urn_secret: str | None = Header(default=None), db: Session = Depends(get_db)):
    urn = _authenticate_urn(db, urn_code, request, x_urn_secret)
    requested = bool(urn.reprint_authorized)
    db.commit()
    return {"requested": requested}


@app.post("/api/urna/{urn_code}/reprint-complete")
def urn_reprint_complete(request: Request, urn_code: str, success: bool = Form(True), x_urn_secret: str | None = Header(default=None), db: Session = Depends(get_db)):
    urn = _authenticate_urn(db, urn_code, request, x_urn_secret)
    if urn.managed_print:
        raise HTTPException(409, "A impressão é controlada pelo aplicativo da urna.")
    urn.reprint_authorized = False
    if success:
        urn.status = UrnStatus.AVAILABLE.value
        urn.pending_reprint_ballot_id = None
    else:
        urn.status = UrnStatus.PRINT_ERROR.value
    db.commit()
    return {"ok": True}


@app.get("/reports/occurrences.pdf")
def report_occurrences(user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    pdf = occurrences_pdf(db)
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": "inline; filename=ocorrencias.pdf"})

@app.get("/reports/logs.pdf")
def report_logs(event_type: str | None = None, user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    pdf = audit_log_pdf(db, event_type=event_type or None)
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": "inline; filename=log-operacional.pdf"})


def _operational_log_text(db: Session, limit: int = 500) -> str:
    election = get_election(db)
    events = db.execute(select(AuditEvent).order_by(AuditEvent.id.desc()).limit(limit)).scalars().all()
    events.reverse()
    lines = [election.institution_name or "INSTITUIÇÃO", election.name, "LOG OPERACIONAL", "="*40]
    for e in events:
        try:
            details = json.loads(e.details_json or "{}")
            detail = ", ".join(f"{k}={v}" for k,v in details.items())
        except Exception:
            detail = ""
        lines.append(f"{e.created_at.strftime('%d/%m/%Y %H:%M:%S')} | {e.event_type} | {e.actor}" + (f" | {detail}" if detail else ""))
    lines += ["="*40, f"Eventos: {len(events)}", f"Versão: {APP_VERSION}"]
    return "\n".join(lines)



@app.post("/api/urna/{urn_code}/printer-config")
def urn_printer_config(
    request: Request, urn_code: str, printer_name: str = Form(""), agent_url: str = Form("http://127.0.0.1:8765"), printer_mode: str = Form("windows"),
    x_urn_secret: str | None = Header(default=None), db: Session = Depends(get_db),
):
    urn = _authenticate_urn(db, urn_code, request, x_urn_secret)
    if urn.managed_print:
        raise HTTPException(409, "A impressão é controlada pelo aplicativo da urna.")
    urn.printer_name = printer_name.strip() or None
    urn.print_agent_url = (agent_url or "http://127.0.0.1:8765").strip()
    urn.printer_mode = printer_mode if printer_mode in ("windows", "escpos") else "windows"
    append_audit(db, "URN_PRINTER_CONFIG_UPDATED", f"URN:{urn.code}", {"printer": urn.printer_name or "", "mode": urn.printer_mode})
    db.commit()
    return {"ok": True, "printer_name": urn.printer_name, "agent_url": urn.print_agent_url, "printer_mode": urn.printer_mode}


@app.get("/api/urna/{urn_code}/device-config")
def urn_device_config(request: Request, urn_code: str, x_urn_secret: str | None = Header(default=None), db: Session = Depends(get_db)):
    urn = _authenticate_urn(db, urn_code, request, x_urn_secret)
    db.commit()
    return {
        "managed_print": bool(urn.managed_print),
        "native_online": deployment.native_online(urn),
        "urn_status": urn.status,
        "printer_ready": bool(urn.printer_ready),
        "printer_name": urn.printer_name or "",
        "agent_url": urn.print_agent_url or "http://127.0.0.1:8765",
        "printer_mode": urn.printer_mode or "escpos",
        "print_enabled": bool(urn.print_enabled and get_election(db).paper_enabled),
        "election_state": get_election(db).state,
        "voting_mode": urn.voting_mode_override if (urn.voting_mode_override and urn.voting_mode_override != "INHERIT") else get_election(db).voting_mode,
        "voting_mode_source": "urn" if (urn.voting_mode_override and urn.voting_mode_override != "INHERIT") else "election",
        "sound_enabled": bool(get_election(db).sound_enabled),
        "keyboard_confirm_key": get_election(db).keyboard_confirm_key,
        "keyboard_correct_key": get_election(db).keyboard_correct_key,
        "keyboard_blank_key": get_election(db).keyboard_blank_key,
    }


@app.post("/api/urna/{urn_code}/print-complete")
def urn_print_complete(request: Request, urn_code: str, ballot_id: str = Form(...), x_urn_secret: str | None = Header(default=None), db: Session = Depends(get_db)):
    urn = _authenticate_urn(db, urn_code, request, x_urn_secret)
    if urn.managed_print:
        raise HTTPException(409, "A impressão é controlada pelo aplicativo da urna.")
    if urn.pending_reprint_ballot_id == ballot_id:
        urn.pending_reprint_ballot_id = None
        urn.reprint_authorized = False
        if urn.status == UrnStatus.PRINT_ERROR.value:
            urn.status = UrnStatus.AVAILABLE.value
    db.commit()
    return {"ok": True}


@app.post("/api/urna/{urn_code}/print-failed")
def urn_print_failed(request: Request, urn_code: str, ballot_id: str = Form(...), x_urn_secret: str | None = Header(default=None), db: Session = Depends(get_db)):
    urn = _authenticate_urn(db, urn_code, request, x_urn_secret)
    if urn.managed_print:
        raise HTTPException(409, "A impressão é controlada pelo aplicativo da urna.")
    if urn.pending_reprint_ballot_id == ballot_id:
        urn.status = UrnStatus.PRINT_ERROR.value
        append_audit(db, "BALLOT_PRINT_FAILED", f"URN:{urn.code}", {"urn": urn.code})
    db.commit()
    return {"ok": True}


def _zero_thermal_text(db: Session, urn: Urn | None = None) -> str:
    election = get_election(db)
    if not election.zero_snapshot_json:
        raise HTTPException(409, "A zerésima ainda não foi emitida.")
    snap = json.loads(election.zero_snapshot_json)
    urn_label = f"{urn.code} - {urn.name}" if urn else "MESA ELEITORAL"
    lines = [
        "ZERÉSIMA",
        election.institution_name or "INSTITUIÇÃO",
        election.name,
        "=" * 42,
        f"Estação: {urn_label}",
        f"Emissão: {snap.get('issued_at','')}",
        f"Versão do sistema: {snap.get('software_version', APP_VERSION)}",
        f"ID da eleição: {snap.get('election_uuid') or election.election_uuid}",
        (f"Eleitores cadastrados: {snap.get('registered_voters',0)}" if snap.get("voter_identification_enabled", True) else "Identificação nominal: DESATIVADA"),
        "-" * 42,
        "CHAPAS REGISTRADAS / VOTOS",
    ]
    for slate in snap.get("slates", []):
        lines.append(f"CHAPA {int(slate['number']):02d} - {slate['name']}")
        lines.append("Votos: 0")
    if snap.get("allow_blank"):
        lines += ["BRANCO", "Votos: 0"]
    lines += [
        "-" * 42,
        "TOTAL DE VOTOS REGISTRADOS: 0",
        f"Hash da configuração: {snap.get('config_hash') or '—'}",
        f"Hash da zerésima: {election.zero_snapshot_hash or '—'}",
        "=" * 42,
        "ASSINATURAS",
        "", "____________________________", "Assinatura 1",
        "", "____________________________", "Assinatura 2",
        "", "____________________________", "Assinatura 3",
        "",
    ]
    return "\n".join(lines)


def _result_thermal_text(db: Session, urn: Urn | None = None) -> str:
    """Boletim textual para a impressora térmica da urna.

    O boletim apresenta somente dados objetivos da eleição e as quantidades
    apuradas. Não calcula colocação, vencedor ou percentuais.
    """
    election = get_election(db)
    if not election.final_tally_json:
        raise HTTPException(409, "A apuração final ainda não foi realizada.")
    data = json.loads(election.final_tally_json)
    registered = len(db.execute(select(Voter)).scalars().all())
    slates = db.execute(select(Slate).where(Slate.active == True).order_by(Slate.number)).scalars().all()
    urn_label = f"{urn.code} - {urn.name}" if urn else "URNA"
    # Cada cabine imprime somente seus próprios totais. O PDF administrativo
    # permanece agregado. Os totais por urna são produzidos durante a apuração
    # local, sem depender de módulo de publicação externo.
    urn_data = (data.get("per_urn") or {}).get(urn.code) if urn is not None else None
    source = urn_data or data
    total = int(source.get("total", 0))
    result_by_number = {int(item.get("number", 0)): int(item.get("votes", 0)) for item in source.get("slates", [])}
    blank_count = int(source.get("blank", 0))
    invalid_count = int(source.get("invalid", 0))
    closed = election.closed_at or datetime.utcnow()
    lines = [
        "BOLETIM DE URNA",
        election.institution_name or "INSTITUIÇÃO",
        election.name,
        "=" * 42,
        f"Urna: {urn_label}",
        f"Data: {closed.strftime('%d/%m/%Y')}",
        f"Abertura: {election.opened_at.strftime('%H:%M:%S') if election.opened_at else '—'}",
        f"Fechamento: {election.closed_at.strftime('%H:%M:%S') if election.closed_at else '—'}",
        f"Versão: {APP_VERSION}",
        "-" * 42,
    ]
    if election.voter_identification_enabled:
        lines += [
            f"Eleitores aptos: {registered:03d}",
            f"Comparecimento: {total:03d}",
            f"Eleitores faltosos: {max(0, registered-total):03d}",
        ]
    else:
        lines += [
            "Identificação nominal: DESATIVADA",
            f"Total de votos: {total:03d}",
        ]
    lines += [
        "=" * 42,
        "CHAPAS",
        "=" * 42,
    ]
    for slate in slates:
        lines.append(f"CHAPA {slate.number:02d} - {slate.name}")
        for member in sorted(slate.members, key=lambda m: (m.display_order, m.id)):
            role = (member.role_name or "INTEGRANTE").upper()
            lines.append(f"{role}: {member.person_name}")
        lines.append(f"VOTOS: {result_by_number.get(slate.number, 0):03d}")
        lines.append("-" * 42)

    # Resumo em três colunas, no padrão do BO oficial:
    # nome à esquerda; número e votos em duas colunas fixas à direita.
    # Mantém 2 dígitos para o número e 3 dígitos para votos.
    name_width = 24
    num_col_width = 8
    votes_col_width = 8

    def summary_name_lines(value: str) -> list[str]:
        text = " ".join(str(value or "").split())
        if not text:
            return [""]
        words = text.split(" ")
        wrapped: list[str] = []
        line = ""
        for word in words:
            while len(word) > name_width:
                if line:
                    wrapped.append(line)
                    line = ""
                wrapped.append(word[:name_width])
                word = word[name_width:]
            candidate = word if not line else line + " " + word
            if len(candidate) <= name_width:
                line = candidate
            else:
                if line:
                    wrapped.append(line)
                line = word
        if line:
            wrapped.append(line)
        return wrapped or [""]

    lines += [
        "=" * 42,
        "RESUMO DAS CHAPAS",
        f"{'NOME DA CHAPA':<{name_width}}"
        f"{'NUM':>{num_col_width}}"
        f"{'VOTOS':>{votes_col_width}}",
        "-" * 42,
    ]
    for slate in slates:
        votes = result_by_number.get(slate.number, 0)
        wrapped = summary_name_lines(slate.name)
        number_text = f"{slate.number:02d}"
        votes_text = f"{votes:03d}"
        lines.append(
            f"{wrapped[0]:<{name_width}}"
            f"{number_text:>{num_col_width}}"
            f"{votes_text:>{votes_col_width}}"
        )
        for continuation in wrapped[1:]:
            lines.append(
                f"{continuation:<{name_width}}"
                f"{'':>{num_col_width}}"
                f"{'':>{votes_col_width}}"
            )
    lines.append("-" * 42)

    if election.allow_blank:
        lines.append(f"BRANCOS: {blank_count:03d}")
    if invalid_count:
        lines.append(f"REGISTROS INVÁLIDOS: {invalid_count:03d}")
    lines += [
        f"TOTAL APURADO: {total:03d}",
        "=" * 42,
        f"ID da eleição: {data.get('election_uuid') or election.election_uuid}",
        f"Hash final dos votos: {data.get('ballot_head_hash') or '—'}",
        f"Hash da configuração: {data.get('config_hash') or '—'}",
        f"Hash da zerésima: {data.get('zero_snapshot_hash') or '—'}",
        "",
    ]
    return "\n".join(lines)




def _zero_thermal_payload(db: Session, urn: Urn | None = None) -> dict:
    election = get_election(db)
    if not election.zero_snapshot_json:
        raise HTTPException(409, "A zerésima ainda não foi emitida.")
    snap = json.loads(election.zero_snapshot_json)
    urn_label = f"{urn.code} - {urn.name}" if urn else "MESA ELEITORAL"
    slates = db.execute(select(Slate).where(Slate.active == True).order_by(Slate.number)).scalars().all()

    def member_payload(slate: Slate) -> list[dict]:
        return [
            {"role": (m.role_name or "Integrante"), "name": m.person_name}
            for m in sorted(slate.members, key=lambda m: (m.display_order, m.id))
        ]

    return {
        "institution": election.institution_name or "INSTITUIÇÃO",
        "election": election.name,
        "station_label": urn_label,
        "issued_at": str(snap.get("issued_at") or ""),
        "version": str(snap.get("software_version") or APP_VERSION),
        "election_uuid": snap.get("election_uuid") or election.election_uuid,
        "registered_voters": int(snap.get("registered_voters") or 0) if snap.get("voter_identification_enabled", True) else None,
        "allow_blank": bool(snap.get("allow_blank")),
        "config_hash": snap.get("config_hash") or election.config_hash or "—",
        "zero_snapshot_hash": election.zero_snapshot_hash or "—",
        "slates": [
            {
                "number": int(s.number),
                "name": s.name,
                "votes": 0,
                "members": member_payload(s),
            }
            for s in slates
        ],
    }


def _result_thermal_payload(db: Session, urn: Urn | None = None) -> dict:
    election = get_election(db)
    if not election.final_tally_json:
        raise HTTPException(409, "A apuração final ainda não foi realizada.")
    data = json.loads(election.final_tally_json)
    registered = len(db.execute(select(Voter)).scalars().all())
    slates = db.execute(select(Slate).where(Slate.active == True).order_by(Slate.number)).scalars().all()
    urn_label = f"{urn.code} - {urn.name}" if urn else "URNA"
    urn_data = (data.get("per_urn") or {}).get(urn.code) if urn is not None else None
    source = urn_data or data
    total = int(source.get("total", 0))
    result_by_number = {int(item.get("number", 0)): int(item.get("votes", 0)) for item in source.get("slates", [])}
    blank_count = int(source.get("blank", 0))

    def member_payload(slate: Slate) -> list[dict]:
        return [
            {"role": (m.role_name or "Integrante"), "name": m.person_name}
            for m in sorted(slate.members, key=lambda m: (m.display_order, m.id))
        ]

    return {
        "institution": election.institution_name or "INSTITUIÇÃO",
        "election": election.name,
        "urn_label": urn_label,
        "date": (election.closed_at.strftime('%d/%m/%Y') if election.closed_at else datetime.utcnow().strftime('%d/%m/%Y')) ,
        "opening": election.opened_at.strftime('%H:%M:%S') if election.opened_at else '—',
        "closing": election.closed_at.strftime('%H:%M:%S') if election.closed_at else '—',
        "version": APP_VERSION,
        "registered": registered,
        "total": total,
        "absent": max(0, registered - total),
        "allow_blank": bool(election.allow_blank),
        "blank_count": blank_count,
        "slates": [
            {
                "number": int(s.number),
                "name": s.name,
                "votes": int(result_by_number.get(s.number, 0)),
                "members": member_payload(s),
            }
            for s in slates
        ],
        "election_uuid": data.get('election_uuid') or election.election_uuid,
        "ballot_head_hash": data.get('ballot_head_hash') or '—',
        "config_hash": data.get('config_hash') or '—',
        "zero_snapshot_hash": data.get('zero_snapshot_hash') or election.zero_snapshot_hash or '—',
    }

def _queue_device_document(db: Session, kind: str, text, copies: int = 1, urn_codes: set[str] | None = None) -> int:
    urns = db.execute(select(Urn).order_by(Urn.code)).scalars().all()
    if urn_codes:
        urns = [u for u in urns if u.code in urn_codes]
    for urn in urns:
        urn_text = text(urn) if callable(text) else text
        layout = "text"
        paper = None
        title = "Urna Escolar"
        if kind == "zero":
            layout = "zero_receipt"
            paper = _zero_thermal_payload(db, urn)
            title = f"Zerésima - {urn.code}"
        elif kind == "result":
            layout = "result_receipt"
            paper = _result_thermal_payload(db, urn)
            title = f"Boletim de Urna - {urn.code}"

        payload = json.dumps({
            "text": urn_text,
            "copies": max(1, min(copies, 20)),
            "cut": True,
            "title": title,
            "layout": layout,
            "paper": paper,
            "mode": urn.printer_mode or "windows",
            "paper_width_mm": 80,
        }, ensure_ascii=False)
        db.add(DevicePrintCommand(urn_id=urn.id, kind=kind, payload_json=payload))
    return len(urns)



@app.get("/api/urna/{urn_code}/print-command")
def urn_print_command(request: Request, urn_code: str, x_urn_secret: str | None = Header(default=None), db: Session = Depends(get_db)):
    urn = _authenticate_urn(db, urn_code, request, x_urn_secret)
    if urn.managed_print:
        raise HTTPException(409, "A impressão é controlada pelo aplicativo da urna.")
    cmd = db.execute(
        select(DevicePrintCommand).where(DevicePrintCommand.urn_id == urn.id, DevicePrintCommand.completed_at.is_(None)).order_by(DevicePrintCommand.created_at).limit(1)
    ).scalar_one_or_none()
    db.commit()
    if not cmd:
        return {"pending": False}
    return {"pending": True, "id": cmd.id, "kind": cmd.kind, "payload": json.loads(cmd.payload_json)}


@app.post("/api/urna/{urn_code}/print-command/{command_id}/complete")
def urn_print_command_complete(
    request: Request, urn_code: str, command_id: str, success: bool = Form(True), x_urn_secret: str | None = Header(default=None), db: Session = Depends(get_db)
):
    urn = _authenticate_urn(db, urn_code, request, x_urn_secret)
    if urn.managed_print:
        raise HTTPException(409, "A impressão é controlada pelo aplicativo da urna.")
    cmd = db.get(DevicePrintCommand, command_id)
    if not cmd or cmd.urn_id != urn.id:
        raise HTTPException(404, "Comando não encontrado.")
    cmd.completed_at = datetime.utcnow()
    cmd.success = bool(success)
    if not success and cmd.kind in ("zero", "result"):
        urn.status = UrnStatus.PRINT_ERROR.value
    append_audit(db, "URN_DOCUMENT_PRINT_OK" if success else "URN_DOCUMENT_PRINT_FAILED", f"URN:{urn.code}", {"document": cmd.kind})
    db.commit()
    return {"ok": True}



def _preflight(db: Session):
    election = get_election(db)
    voters = db.execute(select(Voter)).scalars().all()
    slates = db.execute(select(Slate).where(Slate.active == True)).scalars().all()
    urns = db.execute(select(Urn).order_by(Urn.code)).scalars().all()
    mesarios = db.execute(select(User).where(User.active == True, User.role == UserRole.MESARIO.value)).scalars().all()
    checks = [
        {"key":"voters", "label":"Eleitores cadastrados", "detail":(f"{len(voters)} eleitor(es)" if election.voter_identification_enabled else "Identificação nominal desativada"), "ok":((len(voters)>0) if election.voter_identification_enabled else True), "blocking":True},
        {"key":"slates", "label":"Chapas cadastradas", "detail":f"{len(slates)} chapa(s)", "ok":len(slates)>0, "blocking":True},
        {"key":"urns", "label":"Urnas configuradas", "detail":f"{len(urns)} de {MAX_URNS}", "ok":0<len(urns)<=MAX_URNS, "blocking":True},
        {"key":"mesarios", "label":"Mesário cadastrado", "detail":f"{len(mesarios)} usuário(s) MESÁRIO", "ok":len(mesarios)>0, "blocking":True},
        {"key":"keys", "label":"Chave interna da eleição", "detail":"Gerada e guardada no servidor" if (PUBLIC_KEY_PATH.exists() and PRIVATE_KEY_PATH.exists() and PRIVATE_KEY_PASSWORD_PATH.exists()) else "Será gerada automaticamente na lacração", "ok":((PUBLIC_KEY_PATH.exists() and PRIVATE_KEY_PATH.exists() and PRIVATE_KEY_PASSWORD_PATH.exists()) or election.state == ElectionState.CONFIG.value), "blocking":True},
        {"key":"zero", "label":"Zerésima emitida", "detail": election.zero_issued_at.strftime('%d/%m/%Y %H:%M') if election.zero_issued_at else "Ainda não emitida", "ok":bool(election.zero_issued_at), "blocking": election.state == ElectionState.SEALED.value},
    ]
    if election.state != ElectionState.CONFIG.value:
        try:
            verify_sealed_integrity(db)
            sealed_ok, sealed_detail = True, "Configuração, software e chave interna correspondem à lacração"
        except DomainError as exc:
            sealed_ok, sealed_detail = False, str(exc)
        audit_ok, _ = verify_audit_chain(db)
        checks += [
            {"key":"seal_integrity", "label":"Integridade da lacração", "detail":sealed_detail, "ok":sealed_ok, "blocking":True},
            {"key":"audit_integrity", "label":"Cadeia de auditoria", "detail":"Íntegra" if audit_ok else "Inconsistente", "ok":audit_ok, "blocking":True},
        ]
    if election.paper_enabled:
        zero_ok = 0
        if urns and election.zero_issued_at:
            for u in urns:
                last_zero = db.execute(select(DevicePrintCommand).where(DevicePrintCommand.urn_id == u.id, DevicePrintCommand.kind == "zero").order_by(DevicePrintCommand.created_at.desc()).limit(1)).scalar_one_or_none()
                if last_zero and last_zero.success is True:
                    zero_ok += 1
        configured = sum(1 for u in urns if u.printer_name)
        online = sum(1 for u in urns if deployment.native_online(u) if u.managed_print) + sum(1 for u in urns if not u.managed_print and u.agent_online)
        ready = sum(1 for u in urns if u.printer_ready and (not u.managed_print or deployment.native_online(u)))
        checks += [
            {"key":"printers", "label":"Impressoras das urnas", "detail":f"{configured}/{len(urns)} selecionada(s)", "ok":bool(urns) and configured==len(urns), "blocking":True},
            {"key":"agents", "label":"Aplicativos das urnas", "detail":f"{online}/{len(urns)} online", "ok":bool(urns) and online==len(urns), "blocking":True},
            {"key":"printer_ready", "label":"Impressoras prontas", "detail":f"{ready}/{len(urns)} prontas", "ok":bool(urns) and ready==len(urns), "blocking":True},
            {"key":"zero_printed", "label":"Zerésima nas urnas", "detail":f"{zero_ok}/{len(urns)} impressão(ões) confirmada(s)", "ok":bool(urns) and zero_ok==len(urns), "blocking": election.state == ElectionState.SEALED.value},
        ]
    return checks


@app.get("/api/admin/diagnostics")
def diagnostics(user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    election = get_election(db)
    urns = db.execute(select(Urn).order_by(Urn.code)).scalars().all()
    return {"version":APP_VERSION,"state":election.state,"checks":_preflight(db),"security":security_status(db),"urns":[{"code":u.code,"name":u.name,"status":u.status,"last_seen_at":u.last_seen_at.isoformat() if u.last_seen_at else None,"agent_online":bool(u.agent_online),"printer_ready":bool(u.printer_ready),"printer_name":u.printer_name or ""} for u in urns]}


@app.get("/api/admin/integrity")
def admin_integrity(user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    election = get_election(db)
    status = security_status(db)
    live_config = sha256_text(canonical_json(election_config_payload(db)))
    status.update({
        "election_uuid": election.election_uuid,
        "state": election.state,
        "config_current": live_config,
        "config_ok": not election.config_hash or live_config == election.config_hash,
        "zero_snapshot_hash": election.zero_snapshot_hash,
    })
    return status


@app.get("/admin/security-receipt.json")
def security_receipt(user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    election = get_election(db)
    status = security_status(db)
    live_config = sha256_text(canonical_json(election_config_payload(db)))
    receipt = {
        "format": "urna-escolar-security-receipt-v1",
        "generated_at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
        "app_version": APP_VERSION,
        "election_uuid": election.election_uuid,
        "election_name": election.name,
        "institution_name": election.institution_name,
        "state": election.state,
        "sealed_at": election.sealed_at.isoformat() + "Z" if election.sealed_at else None,
        "config_hash_sealed": election.config_hash,
        "config_hash_current": live_config,
        "software_hash_sealed": election.software_hash,
        "software_hash_current": status["software_current"],
        "public_key_fingerprint_sealed": election.public_key_fingerprint,
        "public_key_fingerprint_current": status["public_key_current"],
        "zero_snapshot_hash": election.zero_snapshot_hash,
        "ballot_head_hash": status["ballot_head"],
        "audit_head_hash": status["audit_head"],
        "audit_chain_ok": status["audit_ok"],
    }
    response = JSONResponse(receipt)
    response.headers["Content-Disposition"] = f'attachment; filename="integridade-{election.election_uuid}.json"'
    return response


@app.post("/api/urna/{urn_code}/device-health")
def urn_device_health(request: Request, urn_code: str, agent_online: bool = Form(False), printer_ready: bool = Form(False), x_urn_secret: str | None = Header(default=None), db: Session = Depends(get_db)):
    urn = _authenticate_urn(db, urn_code, request, x_urn_secret)
    if urn.managed_print:
        raise HTTPException(409, "A impressão é controlada pelo aplicativo da urna.")
    urn.agent_online = bool(agent_online)
    urn.printer_ready = bool(printer_ready)
    db.commit()
    return {"ok":True}



@app.get("/reports/rdv.json")
def report_rdv(user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    election = get_election(db)
    if election.state != ElectionState.CLOSED.value:
        raise HTTPException(409, "O RDV só é disponibilizado depois do encerramento.")
    ballots = db.execute(select(Ballot).order_by(Ballot.rdv_slot, Ballot.id)).scalars().all()
    payload = {
        "format": "urna-escolar-rdv-v1", "election_uuid": election.election_uuid, "version": APP_VERSION,
        "records": [{"slot": b.rdv_slot, "ciphertext": b.ciphertext, "urn": b.urn_code} for b in ballots],
    }
    response = JSONResponse(payload)
    response.headers["Content-Disposition"] = f'attachment; filename="rdv-{election.election_uuid}.json"'
    return response


@app.post("/admin/new-election")
def new_election(confirm_text: str = Form(...), user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    election = get_election(db)
    if election.state != ElectionState.CLOSED.value:
        raise HTTPException(409, "A eleição atual precisa estar encerrada antes de criar uma nova.")
    if confirm_text.strip().upper() != "NOVA ELEICAO":
        raise HTTPException(400, "Digite NOVA ELEICAO para confirmar.")
    try:
        verify_final_tally_integrity(db)
    except DomainError as exc:
        raise HTTPException(409, "A eleição não pode ser arquivada porque o resultado final falhou na verificação de integridade: " + str(exc)) from exc
    archive_dir = DATA_DIR / "archives"
    archive_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    archive_path = archive_dir / f"eleicao-{election.id}-{stamp}.db"
    # Backup consistente do SQLite antes de limpar a eleição ativa.
    db.commit()
    src = sqlite3.connect(str(DB_PATH)); dst = sqlite3.connect(str(archive_path))
    try: src.backup(dst)
    finally: dst.close(); src.close()
    archive = ElectionArchive(
        name=election.name, institution_name=election.institution_name, closed_at=election.closed_at,
        final_tally_json=election.final_tally_json, config_hash=election.config_hash,
        software_hash=election.software_hash, public_key_fingerprint=election.public_key_fingerprint,
        election_uuid=election.election_uuid, sealed_at=election.sealed_at,
        zero_issued_at=election.zero_issued_at, zero_snapshot_json=election.zero_snapshot_json,
        zero_snapshot_hash=election.zero_snapshot_hash,
        voter_identification_enabled=bool(election.voter_identification_enabled),
        voter_lookup_mode=election.voter_lookup_mode,
        database_file=str(archive_path.relative_to(DATA_DIR)),
    )
    db.add(archive); db.flush()
    # Remove somente dados da eleição; usuários e configurações da estação permanecem.
    db.execute(__import__('sqlalchemy').text("DROP TRIGGER IF EXISTS ballots_no_delete"))
    db.execute(__import__('sqlalchemy').text("DROP TRIGGER IF EXISTS audit_no_delete"))
    db.execute(__import__('sqlalchemy').text("DROP TRIGGER IF EXISTS voter_amendments_no_delete"))
    db.execute(__import__('sqlalchemy').text("DROP TRIGGER IF EXISTS election_config_hash_immutable"))
    db.execute(__import__('sqlalchemy').text("DROP TRIGGER IF EXISTS election_software_hash_immutable"))
    db.execute(__import__('sqlalchemy').text("DROP TRIGGER IF EXISTS election_public_key_hash_immutable"))
    db.execute(__import__('sqlalchemy').text("DROP TRIGGER IF EXISTS election_zero_hash_immutable"))
    db.execute(__import__('sqlalchemy').text("DROP TRIGGER IF EXISTS election_final_tally_immutable"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM device_print_commands"))
    if db.execute(__import__('sqlalchemy').text("SELECT name FROM sqlite_master WHERE type='table' AND name='desk_print_commands'")).first():
        db.execute(__import__('sqlalchemy').text("DELETE FROM desk_print_commands"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM voter_exception_permits"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM urn_device_sessions"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM voting_authorizations"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM anonymous_voting_authorizations"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM ballots"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM slate_members"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM slates"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM voter_amendments"))
    db.execute(__import__('sqlalchemy').text("UPDATE voters SET active_urn_id=NULL"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM voters"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM native_vote_states"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM browser_launches"))
    db.execute(__import__('sqlalchemy').text("UPDATE urns SET status='AVAILABLE', pending_reprint_ballot_id=NULL, reprint_authorized=0, last_vote_at=NULL"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM occurrences"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM audit_events"))
    db.execute(__import__('sqlalchemy').text("DELETE FROM settings WHERE key IN ('ballot_head_hash','final_result_auto_queued','mesario_open_ticket_hash','mesario_open_ticket_user','mesario_open_ticket_expires','mesario_close_ticket_hash','mesario_close_ticket_user','mesario_close_ticket_expires') OR key LIKE 'paper_change_ack:%'"))
    try:
        db.expunge(election)
    except Exception:
        pass
    db.execute(__import__('sqlalchemy').text("DELETE FROM elections"))
    db.add(Election())
    db.execute(__import__('sqlalchemy').text("CREATE TRIGGER IF NOT EXISTS ballots_no_delete BEFORE DELETE ON ballots BEGIN SELECT RAISE(ABORT, 'Ballots are append-only'); END;"))
    db.execute(__import__('sqlalchemy').text("CREATE TRIGGER IF NOT EXISTS audit_no_delete BEFORE DELETE ON audit_events BEGIN SELECT RAISE(ABORT, 'Audit log is append-only'); END;"))
    db.execute(__import__('sqlalchemy').text("CREATE TRIGGER IF NOT EXISTS voter_amendments_no_delete BEFORE DELETE ON voter_amendments BEGIN SELECT RAISE(ABORT, 'Voter amendments are append-only'); END;"))
    append_audit(db, "NEW_ELECTION_CREATED", user.username, {"archive": str(archive_path.name)})
    db.commit()
    # Reinstala todos os gatilhos de integridade removidos apenas para a rotação controlada.
    init_db()
    # Uma nova eleição recebe automaticamente um novo par de chaves internas.
    key_dir = DATA_DIR / "keys"
    if key_dir.exists():
        old_dir = archive_dir / f"keys-{stamp}"
        old_dir.mkdir(exist_ok=True)
        for f in key_dir.glob("election_*"):
            shutil.copy2(f, old_dir / f.name)
            f.unlink(missing_ok=True)
    return RedirectResponse("/admin?section=overview", status_code=303)

@app.post("/admin/final-sound")
async def upload_final_sound(sound_file: UploadFile = File(...), user: User = Depends(require_role(UserRole.ADMIN.value)), db: Session = Depends(get_db)):
    if get_election(db).state != ElectionState.CONFIG.value:
        raise HTTPException(409, "O som de finalização só pode ser alterado antes da lacração.")
    filename = (sound_file.filename or "").lower()
    if not filename.endswith(".mp3"):
        raise HTTPException(400, "Envie um arquivo MP3.")
    content = await sound_file.read(5 * 1024 * 1024 + 1)
    if not content or len(content) > 5 * 1024 * 1024:
        raise HTTPException(400, "O MP3 deve ter entre 1 byte e 5 MB.")
    sounds = USER_STATIC_DIR / "sounds"
    sounds.mkdir(parents=True, exist_ok=True)
    (sounds / "final.mp3").write_bytes(content)
    append_audit(db, "FINAL_SOUND_UPDATED", user.username, {"bytes": len(content)})
    db.commit()
    return RedirectResponse("/admin?section=election", status_code=303)


@app.get("/health")
def health(db: Session = Depends(get_db)):
    election = get_election(db)
    return {"ok": True, "version": APP_VERSION, "state": election.state}


from . import deployment
deployment.register(app, templates)
